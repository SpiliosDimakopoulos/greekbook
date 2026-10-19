"""CLI paths that need monkeypatching: watch, serve, open, and error branches."""

import zipfile
from pathlib import Path
from typing import Any

import docx
import pytest

import greekbook.cli as cli
import greekbook.server as server
from greekbook.cli import main
from greekbook.docx_import import import_docx

from .conftest import make_project


def test_build_variants(novel: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    opened: list[Path] = []
    monkeypatch.setattr(cli, "_open", opened.append)
    out = tmp_path / "out.pdf"
    assert main(["build", str(novel / "book.yaml"), "-o", str(out), "--open", "--keep-typst"]) == 0
    assert opened == [out] and out.with_suffix(".typ").is_file()
    assert main(["build", str(novel), "--typst-only", "-o", str(tmp_path / "x.typ")]) == 0


def test_build_reports_markdown_warnings(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    make_project(tmp_path, "title: T\nauthor: A\n", {"chapters/01.md": "<div>x</div>\n"})
    assert main(["build", str(tmp_path)]) == 0
    assert "HTML" in capsys.readouterr().out
    assert main(["validate", str(tmp_path), "--strict"]) == 1


def test_watch_builds_then_stops(novel: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    calls = {"n": 0}

    def fake_sleep(_: float) -> None:
        calls["n"] += 1
        if calls["n"] == 1:
            (novel / "book.yaml").write_text("title: [", encoding="utf-8")  # broken → error path
            return
        raise KeyboardInterrupt

    monkeypatch.setattr(cli.time, "sleep", fake_sleep)
    assert main(["watch", str(novel)]) == 0
    assert (novel / "novel.pdf").is_file()


def test_serve_command(novel: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    seen: dict[str, Any] = {}

    def fake_serve(root: Path, port: int, open_browser: bool) -> None:
        seen.update(root=root, port=port, open_browser=open_browser)

    monkeypatch.setattr(server, "serve", fake_serve)
    assert main(["serve", str(novel), "--port", "9999", "--no-browser"]) == 0
    assert seen == {"root": novel, "port": 9999, "open_browser": False}

    def failing_serve(*_: Any, **__: Any) -> None:
        raise OSError("address in use")

    monkeypatch.setattr(server, "serve", failing_serve)
    assert main(["serve", str(novel)]) == 1


def test_serve_bad_project(tmp_path: Path) -> None:
    assert main(["serve", str(tmp_path)]) == 1
    assert main(["validate", str(tmp_path)]) == 1


def test_serve_function_starts_and_stops(novel: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    def interrupt(self: Any) -> None:
        raise KeyboardInterrupt

    monkeypatch.setattr(server.ThreadingHTTPServer, "serve_forever", interrupt)
    server.serve(novel, port=0, open_browser=False)


def test_import_default_directory(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    d = docx.Document()
    d.add_paragraph("Μόνο κείμενο.")
    src = tmp_path / "simple.docx"
    d.save(str(src))
    monkeypatch.chdir(tmp_path)
    assert main(["import", str(src)]) == 0  # warns: no headings
    assert (tmp_path / "simple" / "book.yaml").is_file()
    assert main(["import", str(src)]) == 1  # target not empty


FOOTNOTES_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    '<w:footnote w:id="1"><w:p><w:r><w:t>Πηγή: ΕΛΣΤΑΤ.</w:t></w:r></w:p></w:footnote>'
    "</w:footnotes>"
)


def _add_footnotes(path: Path) -> None:
    """python-docx cannot write footnotes, so patch the package by hand."""
    with zipfile.ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist()}
    files["word/footnotes.xml"] = FOOTNOTES_XML.encode()
    rels = files["word/_rels/document.xml.rels"].decode()
    files["word/_rels/document.xml.rels"] = rels.replace(
        "</Relationships>",
        '<Relationship Id="rIdFn" Type="http://schemas.openxmlformats.org/officeDocument/'
        '2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>',
    ).encode()
    types = files["[Content_Types].xml"].decode()
    files["[Content_Types].xml"] = types.replace(
        "</Types>",
        '<Override PartName="/word/footnotes.xml" ContentType="application/'
        'vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>',
    ).encode()
    body = files["word/document.xml"].decode()
    body = body.replace(
        "<w:t>ΣΗΜΕΙΩΣΗ</w:t></w:r>",
        '<w:t>ΣΗΜΕΙΩΣΗ</w:t></w:r><w:r><w:footnoteReference w:id="1"/></w:r>'
        '<w:hyperlink r:id="rIdLink"><w:r><w:t>σύνδεσμος</w:t></w:r></w:hyperlink>'
        '<w:r><w:br w:type="page"/></w:r><w:r><w:tab/><w:t>τέλος</w:t></w:r>',
    )
    files["word/document.xml"] = body.encode()
    files["word/_rels/document.xml.rels"] = (
        files["word/_rels/document.xml.rels"]
        .decode()
        .replace(
            "</Relationships>",
            '<Relationship Id="rIdLink" Type="http://schemas.openxmlformats.org/officeDocument/'
            '2006/relationships/hyperlink" Target="https://typst.app" TargetMode="External"/>'
            "</Relationships>",
        )
        .encode()
    )
    with zipfile.ZipFile(path, "w") as z:
        for name, data in files.items():
            z.writestr(name, data)


def test_import_footnotes_links_breaks(tmp_path: Path) -> None:
    d = docx.Document()
    d.add_heading("Κεφάλαιο", 1)
    d.add_paragraph("ΣΗΜΕΙΩΣΗ")
    d.add_paragraph("Επικεφαλίδα με outline", style="Heading 3")
    src = tmp_path / "fn.docx"
    d.save(str(src))
    _add_footnotes(src)
    target = tmp_path / "out"
    import_docx(src, target)
    text = (target / "chapters" / "01-kefalaio.md").read_text(encoding="utf-8")
    assert "ΣΗΜΕΙΩΣΗ^[Πηγή: ΕΛΣΤΑΤ.]" in text
    assert "[σύνδεσμος](https://typst.app)" in text
    assert "\\newpage" in text and " τέλος" in text
    assert "### Επικεφαλίδα με outline" in text
    assert main(["build", str(target)]) == 0
