"""Project creation, the terminal menu, Greek CLI messages and friendly errors."""

import builtins
import time
from pathlib import Path
from typing import Any

import docx
import pytest

import greekbook.cli as cli
import greekbook.server as server
from greekbook.build import BuildError, build, generate_typst, humanize_typst_error
from greekbook.cli import main
from greekbook.config import load_config
from greekbook.lint import lint_project
from greekbook.projects import (
    create_project,
    default_workspace,
    list_projects,
    remember,
    slugify,
    unique_folder,
)


def test_slugify() -> None:
    assert slugify("Η Εργασία μου!") == "i-ergasia-mou"
    assert slugify("Ευχαριστίες") == "evcharisties"
    assert slugify("???") == "ergasia"
    assert slugify("Κεφάλαιο Ψ", "x") == "kefalaio-ps"


@pytest.mark.parametrize("kind", ["thesis", "book", "notes"])
@pytest.mark.parametrize("example", [False, True])
def test_create_project_builds_cleanly(tmp_path: Path, kind: str, example: bool) -> None:
    target = create_project(
        tmp_path / "p", kind, "Τίτλος", "Όνομα", example=example,
        thesis={"institution": "ΕΚΠΑ", "supervisor": ""},
    )  # fmt: skip
    cfg = load_config(target)
    assert cfg.title == "Τίτλος" and cfg.author == "Όνομα"
    if kind == "thesis":
        assert cfg.thesis.institution == "ΕΚΠΑ"
    assert [i for i in lint_project(cfg) if i.level == "error"] == []
    assert build(cfg).output is not None
    assert list_projects()[0].title == "Τίτλος"


def test_create_project_refuses_bad_input(tmp_path: Path) -> None:
    (tmp_path / "x.txt").write_text("hi")
    with pytest.raises(FileExistsError):
        create_project(tmp_path, "book", "T", "A")
    with pytest.raises(ValueError):
        create_project(tmp_path / "new", "poem", "T", "A")


def test_workspace_and_recent(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    assert default_workspace().name == "workspace"  # set by the isolating fixture
    monkeypatch.delenv("GREEKBOOK_HOME")
    assert default_workspace().name == "greekbook"
    create_project(tmp_path / "a", "book", "Α", "Β")
    (tmp_path / "gone").mkdir()
    remember(tmp_path / "gone")  # not a project: skipped in the list
    assert [i.title for i in list_projects(tmp_path)] == ["Α"]
    assert unique_folder(tmp_path, "a").name == "a-2"


def test_bibliography_only_when_cited(tmp_path: Path) -> None:
    target = create_project(tmp_path / "p", "thesis", "Τ", "Α")
    assert "#bibliography(" not in generate_typst(load_config(target))[0]
    chapter = target / "chapters" / "01-eisagogi.md"
    chapter.write_text(chapter.read_text(encoding="utf-8") + "\nΌπως [@knuth1984].\n", "utf-8")
    assert "#bibliography(" in generate_typst(load_config(target))[0]


def test_humanize_typst_error() -> None:
    assert "«a.png»" in humanize_typst_error("file not found (searched at /x/figures/a.png)")
    assert "@fig:none" in humanize_typst_error("label `<fig:none>` does not exist in the document")
    assert "@ghost" in humanize_typst_error("key `ghost` does not exist in the bibliography")
    assert "εικόνα" in humanize_typst_error("failed to decode image")
    assert humanize_typst_error("something odd").startswith("Πρόβλημα στη στοιχειοθεσία")


def test_missing_image_error_is_friendly(tmp_path: Path) -> None:
    target = create_project(tmp_path / "p", "book", "Τ", "Α")
    (target / "chapters" / "01-kefalaio.md").write_text("# Α\n\n![x](nope.png)\n", "utf-8")
    with pytest.raises(BuildError, match=r"Δεν βρέθηκε το αρχείο «nope\.png»"):
        build(load_config(target))


# -- terminal ------------------------------------------------------------------------


def feed(monkeypatch: pytest.MonkeyPatch, answers: list[str]) -> None:
    it = iter(answers)
    monkeypatch.setattr(builtins, "input", lambda _="": next(it))


def no_app(monkeypatch: pytest.MonkeyPatch) -> list[Any]:
    calls: list[Any] = []
    monkeypatch.setattr(server, "run_app", lambda project=None, *a, **k: calls.append(project))
    return calls


def test_menu_create_thesis(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    calls = no_app(monkeypatch)
    feed(monkeypatch, ["2", "Η πτυχιακή μου", "Νίκος", "ΑΠΘ", "Τμήμα Χημείας", "", "ν"])
    assert cli.interactive_menu() == 0
    project = calls[0]
    assert load_config(project).thesis.institution == "ΑΠΘ"
    assert (project / "figures" / "routes.svg").is_file()  # the full example was chosen
    assert "δημιουργήθηκε" in capsys.readouterr().out


def test_menu_other_choices(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    calls = no_app(monkeypatch)
    feed(monkeypatch, ["1"])
    assert cli.interactive_menu() == 0 and calls == [None]
    feed(monkeypatch, ["0"])
    assert cli.interactive_menu() == 0
    feed(monkeypatch, ["6"])
    assert cli.interactive_menu() == 0 and "εντολές" in capsys.readouterr().out
    feed(monkeypatch, ["9"])
    assert cli.interactive_menu() == 1
    feed(monkeypatch, ["3", ""])  # no title
    assert cli.interactive_menu() == 1
    feed(monkeypatch, ["5", str(tmp_path / "missing.docx")])
    assert cli.interactive_menu() == 1
    feed(monkeypatch, ["5", ""])
    assert cli.interactive_menu() == 1
    bad = tmp_path / "bad.docx"
    bad.write_text("not word")
    feed(monkeypatch, ["5", str(bad)])
    assert cli.interactive_menu() == 1


def test_menu_import(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    d = docx.Document()
    d.add_heading("Κεφάλαιο", 1)
    d.add_paragraph("Κείμενο.")
    src = tmp_path / "keimeno.docx"
    d.save(str(src))
    calls = no_app(monkeypatch)
    feed(monkeypatch, ["5", f'"{src}"'])  # dragged files arrive in quotes
    assert cli.interactive_menu() == 0
    assert load_config(calls[0]).title


def test_main_without_arguments(monkeypatch: pytest.MonkeyPatch) -> None:
    calls = no_app(monkeypatch)
    monkeypatch.setattr(cli.sys.stdin, "isatty", lambda: False, raising=False)
    assert main([]) == 0 and calls == [None]  # not a terminal → straight to the app
    monkeypatch.setattr(cli, "interactive_menu", lambda: 7)
    monkeypatch.setattr(cli.sys.stdin, "isatty", lambda: True, raising=False)
    assert main([]) == 7


def test_app_command(monkeypatch: pytest.MonkeyPatch, novel: Path, tmp_path: Path) -> None:
    calls = no_app(monkeypatch)
    assert main(["app", str(novel), "--no-browser"]) == 0 and calls == [novel]
    assert main(["app"]) == 0 and calls[-1] is None
    assert main(["app", str(tmp_path)]) == 1


def test_greek_argparse_errors(capsys: pytest.CaptureFixture[str]) -> None:
    with pytest.raises(SystemExit):
        main(["biuld"])
    assert "Μήπως εννοείς «build»" in capsys.readouterr().err
    with pytest.raises(SystemExit):
        main(["new"])
    assert "Λείπει" in capsys.readouterr().err
    with pytest.raises(SystemExit):
        main(["build", "--wat"])
    assert "Δεν αναγνωρίζω" in capsys.readouterr().err
    with pytest.raises(SystemExit):
        main(["xyz123"])
    assert "Διαθέσιμες" in capsys.readouterr().err
    with pytest.raises(SystemExit):
        main(["--help"])
    out = capsys.readouterr().out
    assert "Χρήση:" in out and "Παραδείγματα" in out and "εντολές" in out


def test_unexpected_errors_are_friendly(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    def boom(_: Any) -> int:
        raise RuntimeError("ουπς")

    monkeypatch.setattr(cli, "cmd_templates", boom)
    parser = cli.build_parser
    monkeypatch.setattr(cli, "build_parser", parser)
    assert main(["templates"]) == 1
    err = capsys.readouterr().err
    assert "Κάτι πήγε στραβά: ουπς" in err and "Traceback" not in err
    with pytest.raises(RuntimeError):
        main(["--debug", "templates"])


def test_spinner_runs_in_a_terminal(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    monkeypatch.setattr(cli.sys.stdout, "isatty", lambda: True, raising=False)
    with cli.Spinner("Δοκιμή"):
        time.sleep(0.2)
    assert "Δοκιμή" in capsys.readouterr().out
