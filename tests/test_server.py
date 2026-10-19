import http.client
import json
import threading
import time
from collections.abc import Iterator
from http.server import ThreadingHTTPServer
from pathlib import Path
from typing import Any
from urllib.parse import quote

import docx
import pytest

import greekbook.server as server_mod
from greekbook.server import App, make_handler

TOKEN = "secret-token"
AUTH = {"X-Greekbook-Token": TOKEN}


class Client:
    def __init__(self, port: int, app: App) -> None:
        self.port = port
        self.app = app

    def request(
        self, method: str, path: str, body: bytes | None = None,
        headers: dict[str, str] | None = None,
    ) -> tuple[int, bytes]:  # fmt: skip
        conn = http.client.HTTPConnection("127.0.0.1", self.port, timeout=15)
        conn.request(method, path, body=body, headers=headers or {})
        res = conn.getresponse()
        return res.status, res.read()

    def get_json(self, path: str) -> Any:
        status, body = self.request("GET", path)
        assert status == 200, body
        return json.loads(body)

    def post(
        self, path: str, data: Any = None, raw: bytes | None = None,
        headers: dict[str, str] | None = None,
    ) -> tuple[int, Any]:  # fmt: skip
        body = raw if raw is not None else json.dumps(data or {}).encode()
        status, out = self.request("POST", path, body, {**AUTH, **(headers or {})})
        return status, json.loads(out) if out else None

    def wait_for_build(self, after: int = 0) -> dict[str, Any]:
        for _ in range(300):
            status: dict[str, Any] = self.get_json("/api/status")
            if not status["building"] and status["version"] > after:
                return status
            time.sleep(0.05)
        raise AssertionError("build did not finish")


@pytest.fixture
def start(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Iterator[Any]:
    monkeypatch.setenv("GREEKBOOK_STATE", str(tmp_path / "state.json"))
    servers: list[ThreadingHTTPServer] = []

    def _start(project: Path | None = None) -> Client:
        app = App(tmp_path / "workspace", project)
        port_ref = [0]
        httpd = ThreadingHTTPServer(("127.0.0.1", 0), make_handler(app, TOKEN, port_ref))
        port_ref[0] = httpd.server_address[1]
        threading.Thread(target=httpd.serve_forever, daemon=True).start()
        servers.append(httpd)
        return Client(port_ref[0], app)

    yield _start
    for httpd in servers:
        httpd.shutdown()
        httpd.server_close()


@pytest.fixture
def client(start: Any, novel: Path) -> Client:
    c: Client = start(novel)
    return c


# -- home and projects -------------------------------------------------------------


def test_home_without_project(start: Any) -> None:
    c = start()
    status, body = c.request("GET", "/")
    assert status == 200 and "Τι θα γράψεις σήμερα".encode() in body and TOKEN.encode() in body
    info = c.get_json("/api/app")
    assert info["project"] is None and "thesis" in info["kinds"]
    assert c.request("GET", "/api/files")[0] == 400  # no project open yet
    assert c.request("GET", "/img/novel.png")[0] == 200
    assert c.request("GET", "/img/../server.py")[0] == 404


def test_create_open_close(start: Any) -> None:
    c = start()
    data = {"kind": "thesis", "title": "Η εργασία μου", "author": "Μαρία",
            "thesis": {"institution": "Πανεπιστήμιο Κρήτης"}}  # fmt: skip
    status, out = c.post("/api/projects/create", data)
    assert status == 200 and out["path"].endswith("i-ergasia-mou")
    assert c.wait_for_build()["ok"]
    assert c.get_json("/api/app")["project"]["title"] == "Η εργασία μου"
    assert b"toolbar" in c.request("GET", "/")[1]
    assert c.post("/api/projects/close")[0] == 200
    assert c.get_json("/api/app")["project"] is None
    assert c.post("/api/projects/open", {"path": out["path"]})[0] == 200
    assert c.get_json("/api/app")["projects"][0]["title"] == "Η εργασία μου"


@pytest.mark.parametrize(
    ("data", "message"),
    [({"kind": "poem", "title": "x"}, "Διάλεξε"), ({"kind": "book", "title": "  "}, "τίτλο")],
)
def test_create_errors(start: Any, data: dict[str, str], message: str) -> None:
    status, out = start().post("/api/projects/create", data)
    assert status == 400 and message in out["error"]


def test_open_errors(start: Any, tmp_path: Path) -> None:
    status, out = start().post("/api/projects/open", {"path": str(tmp_path)})
    assert status == 400 and "book.yaml" in out["error"]


def test_import_docx_upload(start: Any, tmp_path: Path) -> None:
    d = docx.Document()
    d.add_heading("Εισαγωγή", 1)
    d.add_paragraph("Κείμενο.")
    path = tmp_path / "ergasia.docx"
    d.save(str(path))
    c = start()
    status, out = c.post("/api/projects/import?kind=notes", raw=path.read_bytes(),
                         headers={"X-Filename": "ergasia.docx"})  # fmt: skip
    assert status == 200 and Path(out["path"]).name == "ergasia"
    assert c.get_json("/api/app")["project"]["template"] == "academic"
    status, out = c.post("/api/projects/import", raw=b"x", headers={"X-Filename": "old.doc"})
    assert status == 400 and ".docx" in out["error"]
    status, out = c.post("/api/projects/import", raw=b"x", headers={"X-Filename": "bad.docx"})
    assert status == 400 and "Word" in out["error"]


# -- editing -----------------------------------------------------------------------


def test_initial_build_and_preview(client: Client) -> None:
    status = client.wait_for_build()
    assert status["ok"] and status["pages"] >= 4 and status["title"] == "Το φως του φάρου"
    code, svg = client.request("GET", "/page/0.svg")
    assert code == 200 and svg.startswith(b"<svg")
    assert client.request("GET", "/page/999.svg")[0] == 404
    code, pdf = client.request("GET", "/book.pdf?download=1")
    assert code == 200 and pdf.startswith(b"%PDF")
    code, epub = client.request("GET", "/book.epub")
    assert code == 200 and epub.startswith(b"PK")
    files = client.get_json("/api/files")
    assert {"path": "chapters/01.md", "group": "chapters", "title": "Η άφιξη"} in files


def test_save_rebuilds(client: Client, novel: Path) -> None:
    before = client.wait_for_build()
    text = '# Νέο\n\nΚείμενο "με" αλλαγή.\n'
    code, _ = client.request("PUT", "/api/file?path=chapters/01.md", text.encode(), AUTH)
    assert code == 200
    assert (novel / "chapters" / "01.md").read_text(encoding="utf-8") == text
    assert client.wait_for_build(before["version"])["ok"]
    assert client.post("/api/build")[0] == 200


def test_settings_roundtrip(client: Client, novel: Path) -> None:
    cfg = client.get_json("/api/config")["config"]
    cfg["subtitle"] = "Νέος υπότιτλος"
    cfg["template"] = "classic"
    code, _ = client.request("PUT", "/api/config", json.dumps(cfg).encode(), AUTH)
    assert code == 200
    assert "Νέος υπότιτλος" in (novel / "book.yaml").read_text(encoding="utf-8")
    cfg["template"] = "fancy"
    code, body = client.request("PUT", "/api/config", json.dumps(cfg).encode(), AUTH)
    assert code == 400 and "template" in json.loads(body)["error"]
    code, _ = client.request("PUT", "/api/config", b"[1]", AUTH)
    assert code == 400


def test_chapter_management(client: Client, novel: Path) -> None:
    _, out = client.post("/api/chapter", {"title": "Επίλογος"})
    assert out["path"] == "chapters/03-epilogos.md"
    _, out = client.post("/api/chapter/rename", {"path": out["path"], "title": "Τέλος"})
    assert out["path"] == "chapters/03-telos.md"
    assert (novel / out["path"]).read_text(encoding="utf-8").startswith("# Τέλος")
    client.post("/api/chapter/move", {"path": "chapters/03-telos.md", "direction": "up"})
    assert (novel / "chapters" / "02-telos.md").is_file()
    assert (novel / "chapters" / "03.md").is_file()  # the old second chapter moved down
    client.post("/api/chapter/move", {"path": "chapters/01.md", "direction": "up"})  # no-op
    client.post("/api/chapter/delete", {"path": "chapters/02-telos.md"})
    assert not (novel / "chapters" / "02-telos.md").exists()
    assert any(p.name.endswith("02-telos.md") for p in (novel / ".kados").iterdir())
    assert client.request("GET", "/api/file?path=.kados/x.md")[0] == 404
    status, _ = client.post("/api/chapter/rename", {"path": "chapters/01.md", "title": " "})
    assert status == 400


def test_front_matter_and_appendix_lists(start: Any, thesis: Path) -> None:
    c = start(thesis)
    _, out = c.post("/api/chapter", {"title": "Ευχαριστίες", "group": "front"})
    assert out["path"] == "front/evcharisties.md"
    assert "front/evcharisties.md" in (thesis / "book.yaml").read_text(encoding="utf-8")
    c.post("/api/chapter/move", {"path": "front/evcharisties.md", "direction": "up"})
    _, out = c.post("/api/chapter/rename", {"path": "front/evcharisties.md", "title": "Πρόλογος"})
    assert out["path"] == "front/prologos.md"
    c.post("/api/chapter/delete", {"path": "front/prologos.md"})
    assert "prologos" not in (thesis / "book.yaml").read_text(encoding="utf-8")
    _, out = c.post("/api/chapter", {"title": "Κώδικας", "group": "appendices"})
    assert out["path"] == "appendices/kodikas.md"
    keys = [e["key"] for e in c.get_json("/api/bib")]
    assert "dantzig1959" in keys


def test_image_upload(client: Client, novel: Path) -> None:
    status, out = client.post("/api/image?for=chapters/01.md", raw=b"\x89PNG fake",
                              headers={"X-Filename": quote("Ο χάρτης.png")})  # fmt: skip
    assert status == 200
    assert "(../figures/o-chartis.png)" in out["markdown"]
    assert (novel / "figures" / "o-chartis.png").is_file()
    status, out = client.post("/api/image", raw=b"x", headers={"X-Filename": quote("Ο χάρτης.png")})
    assert "(figures/o-chartis-2.png)" in out["markdown"]
    assert client.request("GET", "/api/asset?path=figures/o-chartis.png")[0] == 200
    status, out = client.post("/api/image", raw=b"x", headers={"X-Filename": "virus.exe"})
    assert status == 400 and "PNG" in out["error"]


def test_broken_config_is_reported(client: Client) -> None:
    before = client.wait_for_build()
    client.request("PUT", "/api/file?path=book.yaml", b"title: [", AUTH)
    status = client.wait_for_build(before["version"])
    assert not status["ok"] and "YAML" in str(status["error"])
    assert client.get_json("/api/files") == [
        {"path": "book.yaml", "group": "settings", "title": "book.yaml"}
    ]


# -- security ----------------------------------------------------------------------


def test_writes_need_token(client: Client) -> None:
    assert client.request("PUT", "/api/file?path=chapters/01.md", b"x")[0] == 403
    assert client.request("POST", "/api/build", b"{}", {"X-Greekbook-Token": "wrong"})[0] == 403
    assert client.request("POST", "/api/quit", b"{}")[0] == 403


def test_dns_rebinding_is_refused(client: Client) -> None:
    assert client.request("GET", "/api/files", headers={"Host": "evil.example:80"})[0] == 403


@pytest.mark.parametrize("path", ["../outside.md", "chapters/../../x.md", "book.pdf", "%00.md", ""])
def test_path_traversal_is_refused(client: Client, path: str) -> None:
    assert client.request("GET", f"/api/file?path={path}")[0] == 404
    assert client.request("PUT", f"/api/file?path={path}", b"x", AUTH)[0] == 400


def test_bad_requests(client: Client) -> None:
    assert client.request("POST", "/api/build", b"[1]", AUTH)[0] == 400
    assert client.request("POST", "/api/nope", b"{}", AUTH)[0] == 400
    assert client.request("GET", "/api/nope")[0] == 404


def test_quit_and_run_app(client: Client, tmp_path: Path, novel: Path,
                          monkeypatch: pytest.MonkeyPatch) -> None:  # fmt: skip
    assert client.post("/api/quit")[1] == {"bye": True}
    assert client.app.quit_requested.is_set()

    class QuittingApp(server_mod.App):
        def __init__(self, *args: Any) -> None:
            super().__init__(*args)
            self.quit_requested.set()  # as if «Έξοδος» was pressed right away

    monkeypatch.setattr(server_mod, "App", QuittingApp)
    server_mod.run_app(novel, port=0, open_browser=False, workspace=tmp_path)
