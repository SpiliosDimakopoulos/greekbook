from pathlib import Path

import docx
import pytest
from PIL import Image

from greekbook.cli import main
from greekbook.config import load_config
from greekbook.docx_import import ImportError_, _slug, import_docx


@pytest.fixture
def word_file(tmp_path: Path) -> Path:
    image = tmp_path / "pic.png"
    Image.new("RGB", (40, 20), "teal").save(image)
    d = docx.Document()
    d.core_properties.author = "Μαρία Παπαδοπούλου"
    d.add_heading("Η εργασία μου", 0)
    d.add_paragraph("Κείμενο πριν από το πρώτο κεφάλαιο.")
    d.add_heading("Εισαγωγή", 1)
    p = d.add_paragraph("Αυτό είναι ")
    p.add_run("έντονο").bold = True
    p.add_run(" και ")
    p.add_run("πλάγιο").italic = True
    p.add_run(" με σύμβολα * _ $.")
    d.add_heading("Ενότητα", 2)
    d.add_paragraph("πρώτο", style="List Bullet")
    d.add_paragraph("ένα", style="List Number")
    table = d.add_table(rows=2, cols=2)
    for (r, c), text in {(0, 0): "Α", (0, 1): "Β", (1, 0): "1", (1, 1): "2"}.items():
        table.cell(r, c).text = text
    d.add_picture(str(image))
    d.add_heading("Συμπεράσματα", 1)
    d.add_paragraph("Τέλος.")
    path = tmp_path / "ergasia.docx"
    d.save(str(path))
    return path


def test_import_structure(word_file: Path, tmp_path: Path) -> None:
    target = tmp_path / "project"
    assert import_docx(word_file, target) == []
    names = sorted(p.name for p in (target / "chapters").iterdir())
    assert names == ["01-prologos.md", "02-eisagogi.md", "03-symperasmata.md"]
    intro = (target / "chapters" / "02-eisagogi.md").read_text(encoding="utf-8")
    assert intro.startswith("# Εισαγωγή\n")
    assert "**έντονο**" in intro and "*πλάγιο*" in intro
    assert r"\* \_ \$" in intro
    assert "## Ενότητα" in intro and "- πρώτο" in intro and "1. ένα" in intro
    assert "| Α | Β |" in intro and "![](../figures/image01.png)" in intro
    assert (target / "figures" / "image01.png").is_file()
    cfg = load_config(target)
    assert cfg.title == "Η εργασία μου" and cfg.author == "Μαρία Παπαδοπούλου"


def test_imported_project_builds(word_file: Path, tmp_path: Path) -> None:
    target = tmp_path / "p"
    assert main(["import", str(word_file), str(target), "-t", "academic"]) == 0
    assert main(["validate", str(target)]) == 0
    assert main(["build", str(target)]) == 0


def test_import_errors(tmp_path: Path) -> None:
    bad = tmp_path / "x.docx"
    bad.write_text("not a zip")
    with pytest.raises(ImportError_):
        import_docx(bad, tmp_path / "out")
    assert main(["import", str(bad), str(tmp_path / "out2")]) == 1


def test_slug() -> None:
    assert _slug("Εισαγωγή") == "eisagogi"
    assert _slug("Μελλοντική Έρευνα!") == "mellontiki-erevna"
    assert _slug("???") == "chapter"
